"""
Parameters of a saved automation: the parts of the request that change from run to run.

They go in braces in the request: "cadastre {nome}, CPF {cpf}, no TI e salve". The model
plans with the real values (it needs them to plan well); the approved plan is saved with
the placeholders back in place of the values ("fill CPF = {cpf}"), so it works for any
value; and each replay fills the placeholders with the values of that run.
"""

from __future__ import annotations

import csv
import re
from dataclasses import replace
from pathlib import Path

from anchor.planner import Step

PLACEHOLDER = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")


class ParameterError(ValueError):
    """Missing, unknown or empty parameter values."""


def find_parameters(text: str) -> list[str]:
    """The placeholders of a request, in order, without repetitions."""
    seen: list[str] = []
    for name in PLACEHOLDER.findall(text or ""):
        if name not in seen:
            seen.append(name)
    return seen


def check_values(required: list[str], values: dict[str, str]) -> None:
    missing = [n for n in required if not str(values.get(n, "")).strip()]
    unknown = [n for n in values if n not in required]
    if missing:
        raise ParameterError(f"missing value for: {', '.join(missing)}")
    if unknown:
        raise ParameterError(f"unknown parameter: {', '.join(unknown)} (this automation has: "
                             f"{', '.join(required) or 'none'})")


def render(text, values: dict[str, str]):
    """Fills the placeholders of a text (None stays None)."""
    if text is None:
        return None
    return PLACEHOLDER.sub(lambda m: str(values.get(m.group(1), m.group(0))), text)


def render_steps(steps: list[Step], values: dict[str, str]) -> list[Step]:
    return [replace(s, description=render(s.description, values), value=render(s.value, values),
                    expect=render(s.expect, values)) for s in steps]


def templatize(text, values: dict[str, str]):
    """Puts the placeholders back in place of the values (the longest values first)."""
    if text is None:
        return None
    for name, value in sorted(values.items(), key=lambda kv: -len(str(kv[1]))):
        value = str(value)
        if len(value) >= 2:
            text = text.replace(value, "{" + name + "}")
    return text


def templatize_steps(steps: list[Step], values: dict[str, str]) -> list[Step]:
    if not values:
        return list(steps)
    return [replace(s, description=templatize(s.description, values), value=templatize(s.value, values),
                    expect=templatize(s.expect, values)) for s in steps]


def parse_assignments(items: list[str]) -> dict[str, str]:
    """["nome=Maria Silva", "cpf=123"] → {"nome": "Maria Silva", "cpf": "123"}."""
    values = {}
    for item in items or []:
        if "=" not in item:
            raise ParameterError(f"use name=value, not {item!r}")
        name, value = item.split("=", 1)
        values[name.strip()] = value
    return values


def read_rows(path: str | Path) -> list[dict[str, str]]:
    """The rows of a CSV file, one dict per row (the header gives the names). Accepts the
    separator Excel uses in Brazil (;) as well as the comma."""
    return rows_from_text(Path(path).read_text(encoding="utf-8-sig"))


MAX_ROWS = 500           # a spreadsheet run is one queued run: more rows wait too long for a person to watch


def rows_from_file(data: bytes, filename: str) -> list[dict[str, str]]:
    """The rows of a spreadsheet the interface sent: a CSV (UTF-8, or Windows' Latin encoding) or the
    first sheet of an Excel .xlsx, whose first row names the columns. Empty rows are skipped."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        from io import BytesIO

        from openpyxl import load_workbook
        try:
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:                          # not an Excel file, or a damaged one
            raise ParameterError(f"the file could not be read as an Excel workbook ({exc})")
        sheet = book.worksheets[0]
        lines = list(sheet.iter_rows(values_only=True))
        book.close()
        if not lines:
            return []
        header = [cell_text(c) for c in lines[0]]
        rows = [{h: cell_text(v) for h, v in zip(header, line) if h} for line in lines[1:]]
    elif suffix in (".csv", ".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")                  # what Excel on Windows writes as "CSV"
        rows = rows_from_text(text)
    else:
        raise ParameterError("use a .csv or .xlsx file")
    return [r for r in rows if any(v for v in r.values())]


def cell_text(value) -> str:
    """A spreadsheet cell as the text a person would type: 12 rather than 12.0, a date as 2026-10-30."""
    from datetime import date, datetime
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if (value.hour, value.minute, value.second) == (0, 0, 0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def check_rows(parameters: list[str], rows: list[dict[str, str]]) -> dict:
    """What the interface shows before running: the columns the automation needs and the sheet lacks,
    the ones it has and the automation does not use, and the rows with an empty value it needs."""
    columns = list(rows[0].keys()) if rows else []
    missing = [p for p in parameters if p not in columns]
    extra = [c for c in columns if c not in parameters]
    blanks = [n for n, row in enumerate(rows, 1) if any(not row.get(p, "").strip() for p in parameters if p in columns)]
    return {"rows": len(rows), "columns": columns, "missing": missing, "extra": extra, "blank_rows": blanks,
            "too_many": len(rows) > MAX_ROWS}


def rows_from_text(text: str) -> list[dict[str, str]]:
    try:
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.DictReader(text.splitlines(), dialect=dialect))
    return [{k.strip(): (v or "").strip() for k, v in row.items() if k is not None} for row in rows]
